"""Convert architecture Markdown to a readable Word document."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def shade_paragraph(paragraph, fill: str = "F4F4F4") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def strip_md_inline(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    return s


def add_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            part = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", part)
            paragraph.add_run(part)


def convert(src: Path, out: Path) -> None:
    text = src.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for i in range(1, 4):
        hs = doc.styles[f"Heading {i}"]
        hs.font.name = "Calibri"
        hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    doc.add_heading("RasaOps — System Architecture", 0)
    note = doc.add_paragraph()
    nr = note.add_run(
        "Readable Word export of the full architecture design. "
        "Code/diagram blocks are plain text (Mermaid drawings are not rendered as graphics)."
    )
    nr.italic = True
    nr.font.size = Pt(10)
    nr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    meta = doc.add_paragraph()
    meta.add_run("Status: Revised  ·  Date: 2026-07-30  ·  Audience: Engineering, product, pilot ops")

    lines = text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_rows: list[str] = []
    started = False

    def flush_code() -> None:
        nonlocal code_buf
        if not code_buf:
            return
        block = "\n".join(code_buf)
        if len(block) > 6000:
            block = block[:6000] + "\n…(truncated)…"
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(block)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        shade_paragraph(p, "F4F4F4")
        code_buf = []

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        rows: list[list[str]] = []
        for row in table_rows:
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            if cells and all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells):
                continue
            rows.append(cells)
        table_rows = []
        if not rows:
            return
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(cols):
                cell = t.cell(ri, ci)
                val = strip_md_inline(row[ci] if ci < len(row) else "")
                cell.text = val
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Calibri"
                        if ri == 0:
                            run.bold = True
                if ri == 0:
                    set_cell_shading(cell, "D6EAF8")
        doc.add_paragraph()

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                in_code = False
                flush_code()
            else:
                if table_rows:
                    flush_table()
                in_code = True
                lang = line.strip()[3:].strip()
                if lang:
                    code_buf.append(f"[diagram / code: {lang}]")
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if "|" in line and line.strip().startswith("|"):
            table_rows.append(line)
            i += 1
            continue
        if table_rows:
            flush_table()

        if not line.strip() or re.match(r"^---+$", line.strip()):
            i += 1
            continue

        hm = re.match(r"^(#{1,4})\s+(.*)$", line)
        if hm:
            level = len(hm.group(1))
            title_t = strip_md_inline(hm.group(2))
            if level == 1 and not started:
                started = True
                i += 1
                continue
            started = True
            doc.add_heading(title_t, level=min(level, 3))
            i += 1
            continue

        started = True
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(strip_md_inline(line[2:]))
            r.italic = True
            r.font.size = Pt(10)
            i += 1
            continue

        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            content = m.group(3)
            is_num = bool(re.match(r"\d+\.", m.group(2)))
            p = doc.add_paragraph(style="List Number" if is_num else "List Bullet")
            add_runs(p, content)
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1

    if in_code:
        flush_code()
    if table_rows:
        flush_table()

    end = doc.add_paragraph()
    er = end.add_run(
        "End of readable export. Source of truth: docs/architecture/RasaOps-System-Architecture.md"
    )
    er.italic = True
    er.font.size = Pt(9)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"Wrote {out} ({out.stat().st_size} bytes)")


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    src = root / "docs" / "architecture" / "RasaOps-System-Architecture.md"
    out = root / "docs" / "architecture" / "RasaOps-System-Architecture.docx"
    if len(sys.argv) >= 2:
        src = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        out = Path(sys.argv[2])
    if not src.exists():
        print(f"Missing {src}", file=sys.stderr)
        return 1
    convert(src, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
